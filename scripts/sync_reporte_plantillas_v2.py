"""Copia plantillas V2 y aplica encabezados normalizados (MAYUSCULAS, sin acentos)."""
from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl

V2 = Path(r"E:\GeneralDocumentacion\Documentos\GPSolution\ATT\ReportesMonitorSolicitud\V2")
OUT = Path(
    r"E:\Proyectos\GPSolutions\SistemasActualizados\RepositorioGPSoluciones\Ecosistema_ATT"
    r"\API\src\ATT.Monitor.Api\Templates\Reportes"
)

HOJA_LAYOUT = "Nuevo Layout"
HOJA_1 = "Hoja1"

# Bloques de encabezados compartidos entre reportes.
IDENTIFICACION_EP = ["ID", "ESTACION DE PAGO", "REGION", "FORMATO DE TIENDA"]
CRITERIO_PAGO = "CRITERIO DE PAGO"
METODO_PAGO = "METODO DE PAGO"

# Columnas comunes al detalle de transacciones.
TRANSACCION_DATOS = [
    "FECHA", "HORA", "CUENTA", "DN", "TICKET", "NOMBRE CLIENTE", "TIPO DE OPERACION",
    "FORMA DE PAGO", "MONTO PAGADO", "MONTO INGRESADO", "CAMBIO ENTREGADO",
    "CAMBIO PENDIENTE", "ESTATUS",
]
TRANSACCION_CIERRE = ["COD AUTORIZACION", "NUM TARJETA", "REFERENCIA"]

HEADERS: dict[str, tuple[str, list[str]]] = {
    "Reporte de Estaciones de Pago Activas.xlsx": (
        HOJA_LAYOUT,
        [
            *IDENTIFICACION_EP, CRITERIO_PAGO,
            "DIRECCION", "VERSION APLICATIVO", "SERIE BANK", "PSW BANK", "S.O", "HOST NAME",
            "MEMORIA", "SERIE EP", "VERSION TEMPLATE", "IP", "MODELO", "STATUS CAJA",
            "ACEPTADOR", "DISPENSADOR", "IMPRESORA", "PIN PAD", "ESTATUS SW", "CODI", "EFE",
            "TAR", "URL BANK", "URL WEBSERVICE", "FECHA DE ULTIMA ACT",
        ],
    ),
    "Reporte de Transacciones con Error.xlsx": (
        HOJA_LAYOUT,
        [
            *IDENTIFICACION_EP, CRITERIO_PAGO,
            *TRANSACCION_DATOS,
            "CODIGO ERROR",
            *TRANSACCION_CIERRE,
        ],
    ),
    "Reporte de Transacciones por Equipo.xlsx": (
        HOJA_1,
        [
            *IDENTIFICACION_EP, METODO_PAGO,
            *TRANSACCION_DATOS,
            "CANTIDAD REVERSADA",
            *TRANSACCION_CIERRE,
        ],
    ),
    "Reporte de Campañas MKT.xlsx": (
        HOJA_LAYOUT,
        [
            *IDENTIFICACION_EP, CRITERIO_PAGO,
            "NOMBRE CAMPANA", "TIPO", "FECHA INICIO", "FECHA TERMINO", "FECHA CREACION",
            "ESTATUS",
        ],
    ),
    "Reporte de Cierre de Caja.xlsx": (
        HOJA_LAYOUT,
        [
            *IDENTIFICACION_EP, CRITERIO_PAGO,
            "FECHA", "HORA", "REFERENCIA", "BANCO", "CUENTA", "EFECTIVO", "TARJETA", "CODI",
            "CASETERO 1 REM", "CASETERO 1 DISP", "CASETERO 1 RECH",
            "CASETERO 2 REM", "CASETERO 2 DISP", "CASETERO 2 RECH",
            "CASETERO 3 REM", "CASETERO 3 DISP", "CASETERO 3 RECH",
            "MXN20", "MXN50", "MXN100", "MXN200", "MXN500", "MXN1000",
        ],
    ),
    "Reporte de Contadores.xlsx": (
        HOJA_LAYOUT,
        [
            *IDENTIFICACION_EP, CRITERIO_PAGO,
            "DISPENSADO $20", "DISPENSADO $50", "DISPENSADO $100",
            "REMANENTE $20", "REMANENTE $50", "REMANENTE $100",
            "RECHAZADO $20", "RECHAZADO $50", "RECHAZADO $100",
            "ACEPTADOS $20", "ACEPTADOS $50", "ACEPTADOS $100", "ACEPTADOS $200",
            "ACEPTADOS $500", "ACEPTADOS $1000", "EFECTIVO",
        ],
    ),
    "Reporte de Transacciones.xlsx": (
        HOJA_LAYOUT,
        [
            *IDENTIFICACION_EP, CRITERIO_PAGO,
            *TRANSACCION_DATOS,
            "CANTIDAD REVERSADA", "CODIGO ERROR",
            *TRANSACCION_CIERRE,
        ],
    ),
}


def resolver_origen(filename: str) -> Path:
    """Ubica la plantilla V2, tolerando diferencias de codificación en el nombre."""
    src = V2 / filename
    if src.exists():
        return src

    matches = list(V2.glob(filename.replace("ñ", "?")))
    if not matches:
        matches = [p for p in V2.iterdir() if p.name.lower() == filename.lower()]
    if not matches:
        raise FileNotFoundError(f"No se encontro plantilla V2: {filename}")

    return matches[0]


def aplicar_encabezados(dst: Path, filename: str, sheet_name: str, headers: list[str]) -> None:
    wb = openpyxl.load_workbook(dst)
    if sheet_name not in wb.sheetnames:
        raise KeyError(f"{filename}: falta hoja {sheet_name}")

    ws = wb[sheet_name]
    for col, title in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=title)

    # limpiar columnas sobrantes en fila 1
    for col in range(len(headers) + 1, ws.max_column + 1):
        ws.cell(row=1, column=col, value=None)

    wb.save(dst)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for filename, (sheet_name, headers) in HEADERS.items():
        dst = OUT / filename
        shutil.copy2(resolver_origen(filename), dst)
        aplicar_encabezados(dst, filename, sheet_name, headers)
        print(f"OK {filename} ({len(headers)} cols)")


if __name__ == "__main__":
    main()
